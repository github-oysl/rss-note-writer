from __future__ import annotations

import argparse
import fnmatch
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class ColumnMapping:
    plugin_output: str
    id: str
    ip_or_domain: str
    title: str
    severity: str
    related_asset: str


@dataclass(frozen=True)
class AppConfig:
    columns: ColumnMapping
    high_levels: Tuple[str, ...]
    systems: Tuple["SystemConfig", ...]
    iap_titles: Tuple[str, ...]
    iap_subpackage_patterns: Tuple[str, ...]


@dataclass(frozen=True)
class SystemConfig:
    name: str
    ips: Tuple[str, ...]


@dataclass(frozen=True)
class VulnRecord:
    row_index: int
    id: str
    ip_or_domain: str
    title: str
    severity: str
    plugin_output: str
    system: str
    iap_category: str
    service_name: str
    subpackage_name: str


SPRINGBOOT_PATH_RE = re.compile(r"SpringBoot项目包路径：\s*(?P<path>[^\r\n]+)")
SPRINGBOOT_JAR_RE = re.compile(r"SpringBoot子包名称：\s*(?P<jar>[^\r\n]+)")


def load_config(config_path: Path) -> AppConfig:
    """从 JSON 文件加载配置（列映射、系统-IP 列表、IAP 标题列表等）。"""
    raw = json.loads(config_path.read_text(encoding="utf-8"))

    columns_raw = raw.get("columns") or {}
    columns = ColumnMapping(
        plugin_output=str(columns_raw.get("plugin_output") or "插件输出"),
        id=str(columns_raw.get("id") or "ID编号"),
        ip_or_domain=str(columns_raw.get("ip_or_domain") or "IP/域名"),
        title=str(columns_raw.get("title") or "漏洞标题"),
        severity=str(columns_raw.get("severity") or "工单漏洞等级"),
        related_asset=str(columns_raw.get("related_asset") or "关联资产"),
    )

    high_levels = tuple(str(x) for x in (raw.get("high_levels") or ["高危", "超危"]))

    systems_list: List[SystemConfig] = []
    for item in raw.get("systems") or []:
        name = str(item.get("name") or "").strip()
        ips = tuple(str(x).strip() for x in (item.get("ips") or []) if str(x).strip())
        if name:
            systems_list.append(SystemConfig(name=name, ips=ips))

    iap_titles = tuple(str(x).strip() for x in (raw.get("iap_titles") or []) if str(x).strip())
    iap_subpackage_patterns = tuple(
        str(x).strip() for x in (raw.get("iap_subpackage_patterns") or []) if str(x).strip()
    )

    return AppConfig(
        columns=columns,
        high_levels=high_levels,
        systems=tuple(systems_list),
        iap_titles=iap_titles,
        iap_subpackage_patterns=iap_subpackage_patterns,
    )


def build_ip_to_system_map(systems: Sequence[SystemConfig]) -> Dict[str, str]:
    """将系统配置展开为 IP/域名 -> 系统名 的映射表。"""
    mapping: Dict[str, str] = {}
    for system in systems:
        for ip in system.ips:
            key = normalize_ip_or_domain(ip)
            if key and key not in mapping:
                mapping[key] = system.name
    return mapping


def normalize_ip_or_domain(value: str) -> str:
    """规范化 IP/域名字段，便于做系统归属匹配。"""
    s = (value or "").strip()
    if not s:
        return ""
    s = s.split()[0].strip()
    s = s.rstrip("/")
    return s


def is_high_plus(severity: str, high_levels: Sequence[str]) -> bool:
    """判断漏洞等级是否属于“高危以上”。"""
    sev = (severity or "").strip()
    return any(sev == level for level in high_levels)


def match_any_subpackage_pattern(subpackage_name: str, patterns: Sequence[str]) -> bool:
    """判断子包名称是否命中任意通配符模式（支持 * 语法）。"""
    name = (subpackage_name or "").strip()
    if not name:
        return False

    name_lower = name.lower()
    for pattern in patterns:
        p = (pattern or "").strip()
        if not p:
            continue
        if fnmatch.fnmatchcase(name_lower, p.lower()):
            return True
    return False


def classify_iap(title: str, iap_titles: Sequence[str], subpackage_name: str, iap_subpackage_patterns: Sequence[str]) -> str:
    """先按标题匹配 IAP，再按子包通配符匹配 IAP，否则为非IAP。"""
    t = (title or "").strip()
    if t and any(t == x for x in iap_titles):
        return "IAP"
    if match_any_subpackage_pattern(subpackage_name, iap_subpackage_patterns):
        return "IAP"
    return "非IAP"


def dedupe_preserve_order(items: Iterable[str]) -> List[str]:
    """对字符串列表去重并保持首次出现顺序。"""
    seen: Set[str] = set()
    out: List[str] = []
    for item in items:
        key = (item or "").strip()
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(key)
    return out


def parse_plugin_output_pairs(plugin_output: str) -> Tuple[List[Tuple[str, str]], bool]:
    """解析插件输出文本，提取(服务名, 子包名称)对，支持多处匹配并去重。"""
    text = plugin_output or ""

    service_names_raw: List[str] = []
    for m in SPRINGBOOT_PATH_RE.finditer(text):
        path_value = (m.group("path") or "").strip()
        service = extract_service_name_from_path(path_value) or ""
        if service:
            service_names_raw.append(service)

    jar_names_raw: List[str] = []
    for m in SPRINGBOOT_JAR_RE.finditer(text):
        jar = (m.group("jar") or "").strip()
        if jar:
            jar_names_raw.append(jar)

    parsed_any = bool(service_names_raw or jar_names_raw)
    if not parsed_any:
        return [("其他", safe_excel_text(text))], False

    service_names = dedupe_preserve_order(service_names_raw)
    jar_names = dedupe_preserve_order(jar_names_raw)

    if not service_names and jar_names:
        pairs = [("其他", j) for j in jar_names]
        return dedupe_preserve_order_pairs(pairs), True

    if service_names and not jar_names:
        pairs = [(s, "其他") for s in service_names]
        return dedupe_preserve_order_pairs(pairs), True

    if len(service_names) == 1 and len(jar_names) > 1:
        pairs2 = [(service_names[0], safe_excel_text(j)) for j in jar_names]
        return dedupe_preserve_order_pairs(pairs2), True

    if len(jar_names) == 1 and len(service_names) > 1:
        pairs3 = [(s, safe_excel_text(jar_names[0])) for s in service_names]
        return dedupe_preserve_order_pairs(pairs3), True

    max_len = max(len(service_names), len(jar_names))
    pairs4: List[Tuple[str, str]] = []
    for idx in range(max_len):
        s = service_names[idx] if idx < len(service_names) else "其他"
        j = jar_names[idx] if idx < len(jar_names) else "其他"
        pairs4.append((s or "其他", safe_excel_text(j or "其他")))

    return dedupe_preserve_order_pairs(pairs4), True


def dedupe_preserve_order_pairs(pairs: Iterable[Tuple[str, str]]) -> List[Tuple[str, str]]:
    """对(服务名, 子包名称)对去重并保持首次出现顺序。"""
    seen: Set[Tuple[str, str]] = set()
    out: List[Tuple[str, str]] = []
    for service_name, jar_name in pairs:
        s = (service_name or "").strip() or "其他"
        j = (jar_name or "").strip() or "其他"
        key = (s, j)
        if key in seen:
            continue
        seen.add(key)
        out.append(key)
    return out


def extract_service_name_from_path(path_value: str) -> str:
    """从 SpringBoot 项目包路径中提取服务名（按 /home/<服务名>/springboot 规则）。"""
    s = (path_value or "").strip()
    if not s:
        return ""

    marker = "/home/"
    if marker not in s:
        return ""
    after = s.split(marker, 1)[1]
    if "/springboot" not in after:
        return ""
    service = after.split("/springboot", 1)[0].strip("/").strip()
    return service


def safe_excel_text(value: str) -> str:
    """将文本截断到 Excel 单元格可接受的最大长度，避免写入异常。"""
    s = value or ""
    return s[:32767]


def read_vuln_records_from_excel(input_path: Path, columns: ColumnMapping, sheet_name: Optional[str]) -> Tuple[List[Dict[str, Any]], List[str]]:
    """读取 Excel，并返回行字典列表与表头列表。"""
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return [], []

    headers = [str(x).strip() if x is not None else "" for x in header_row]
    header_to_index = {h: i for i, h in enumerate(headers) if h}

    required = [columns.plugin_output, columns.id, columns.ip_or_domain, columns.title, columns.severity]
    missing = [c for c in required if c not in header_to_index]
    if missing:
        raise ValueError(f"Excel 缺少必要列：{', '.join(missing)}")

    rows: List[Dict[str, Any]] = []
    for row in rows_iter:
        if row is None:
            continue
        record: Dict[str, Any] = {}
        for col_name in required:
            idx = header_to_index[col_name]
            record[col_name] = row[idx] if idx < len(row) else None
        related_asset_col = columns.related_asset
        if related_asset_col in header_to_index:
            idx = header_to_index[related_asset_col]
            record[related_asset_col] = row[idx] if idx < len(row) else None
        rows.append(record)

    return rows, headers


def build_records(rows: Iterable[Dict[str, Any]], config: AppConfig, sheet_row_offset: int = 2) -> List[VulnRecord]:
    """将原始行数据转换为带分类字段的漏洞记录列表。"""
    ip_to_system = build_ip_to_system_map(config.systems)
    records: List[VulnRecord] = []

    for i, row in enumerate(rows):
        row_index = i + sheet_row_offset
        vuln_id = str(row.get(config.columns.id) or "").strip()
        ip_or_domain = normalize_ip_or_domain(str(row.get(config.columns.ip_or_domain) or ""))
        title = str(row.get(config.columns.title) or "").strip()
        severity = str(row.get(config.columns.severity) or "").strip()
        plugin_output = str(row.get(config.columns.plugin_output) or "")
        related_asset = str(row.get(config.columns.related_asset) or "").strip()

        system = ip_to_system.get(ip_or_domain, "未归类")
        pairs, _parsed = parse_plugin_output_pairs(plugin_output)
        if not pairs:
            pairs = [("其他", safe_excel_text(plugin_output))]

        for service_name, subpackage_name in pairs:
            final_service_name = related_asset or service_name
            iap_category = classify_iap(
                title=title,
                iap_titles=config.iap_titles,
                subpackage_name=subpackage_name,
                iap_subpackage_patterns=config.iap_subpackage_patterns,
            )
            records.append(
                VulnRecord(
                    row_index=row_index,
                    id=vuln_id,
                    ip_or_domain=ip_or_domain,
                    title=title,
                    severity=severity,
                    plugin_output=safe_excel_text(plugin_output),
                    system=system,
                    iap_category=iap_category,
                    service_name=final_service_name or "其他",
                    subpackage_name=subpackage_name or "其他",
                )
            )

    return records


def summarize_by_system(records: Sequence[VulnRecord], high_levels: Sequence[str]) -> List[Dict[str, Any]]:
    """按系统汇总总漏洞数、高危以上数、标题去重后的漏洞种类数。"""
    bucket: Dict[str, List[VulnRecord]] = {}
    for r in records:
        bucket.setdefault(r.system, []).append(r)

    rows: List[Dict[str, Any]] = []
    for system, group in sorted(bucket.items(), key=lambda x: x[0]):
        total = len({r.row_index for r in group})
        iap_count = len({r.row_index for r in group if r.iap_category == "IAP"})
        non_iap_count = len({r.row_index for r in group if r.iap_category != "IAP"})
        high_plus = len({r.row_index for r in group if is_high_plus(r.severity, high_levels)})
        high_plus_iap = len(
            {r.row_index for r in group if r.iap_category == "IAP" and is_high_plus(r.severity, high_levels)}
        )
        unique_titles = len({r.title for r in group if r.title})
        rows.append(
            {
                "系统": system,
                "总漏洞": total,
                "IAP漏洞数": iap_count,
                "非IAP漏洞数": non_iap_count,
                "高危以上": high_plus,
                "高危以上IAP": high_plus_iap,
                "漏洞种类数(标题去重)": unique_titles,
            }
        )
    return rows


def summarize_by_system_service(records: Sequence[VulnRecord], high_levels: Sequence[str]) -> List[Dict[str, Any]]:
    """按系统 -> 服务名 维度做统计输出，并补充 IAP 子统计。"""
    bucket: Dict[Tuple[str, str], List[VulnRecord]] = {}
    for r in records:
        key = (r.system, r.service_name or "其他")
        bucket.setdefault(key, []).append(r)

    rows: List[Dict[str, Any]] = []
    for (system, service_name), group in sorted(bucket.items(), key=lambda x: x[0]):
        total = len({r.row_index for r in group})
        high_plus_total = len({r.row_index for r in group if is_high_plus(r.severity, high_levels)})
        unique_titles_total = len({r.title for r in group if r.title})

        iap_group = [r for r in group if r.iap_category == "IAP"]
        iap_total = len({r.row_index for r in iap_group})
        high_plus_iap = len({r.row_index for r in iap_group if is_high_plus(r.severity, high_levels)})
        unique_titles_iap = len({r.title for r in iap_group if r.title})

        rows.append(
            {
                "系统": system,
                "服务名": service_name,
                "总漏洞": total,
                "高危以上总漏洞": high_plus_total,
                "漏洞种类数(标题去重)": unique_titles_total,
                "IAP总漏洞": iap_total,
                "高危以上IAP总漏洞": high_plus_iap,
                "IAP漏洞种类数(标题去重)": unique_titles_iap,
            }
        )
    return rows


def summarize_by_system_iap_service_subpackage(records: Sequence[VulnRecord], high_levels: Sequence[str]) -> List[Dict[str, Any]]:
    """按系统 -> 服务名 -> 子包名称 维度做统计输出，并汇总漏洞 ID 与标题。"""
    bucket: Dict[Tuple[str, str, str], List[VulnRecord]] = {}
    for r in records:
        key = (r.system, r.service_name or "其他", r.subpackage_name or "其他")
        bucket.setdefault(key, []).append(r)

    rows: List[Dict[str, Any]] = []
    for (system, service_name, subpackage_name), group in sorted(bucket.items(), key=lambda x: x[0]):
        total = len({r.row_index for r in group})
        high_plus = len({r.row_index for r in group if is_high_plus(r.severity, high_levels)})
        unique_titles = len({r.title for r in group if r.title})
        uniq_ids = dedupe_preserve_order([r.id for r in group if r.id])
        uniq_titles = dedupe_preserve_order([r.title for r in group if r.title])
        rows.append(
            {
                "系统": system,
                "服务名": service_name,
                "子包名称": subpackage_name,
                "总漏洞": total,
                "高危以上": high_plus,
                "漏洞种类数(标题去重)": unique_titles,
                "漏洞Id汇总": ",".join(uniq_ids),
                "漏洞标题汇总": "\n".join(uniq_titles),
            }
        )
    return rows


def write_output_excel(
    output_path: Path,
    records: Sequence[VulnRecord],
    system_summary: Sequence[Dict[str, Any]],
    system_service_summary: Sequence[Dict[str, Any]],
    system_iap_service_subpackage_summary: Sequence[Dict[str, Any]],
) -> None:
    """写出包含明细与统计的 Excel 报表。"""
    wb = Workbook()

    ws_detail = wb.active
    ws_detail.title = "明细"
    detail_headers = [
        "行号",
        "系统",
        "IAP分类",
        "服务名",
        "子包名称",
        "ID编号",
        "IP/域名",
        "漏洞标题",
        "工单漏洞等级",
        "插件输出",
    ]
    ws_detail.append(detail_headers)
    for r in records:
        ws_detail.append(
            [
                r.row_index,
                r.system,
                r.iap_category,
                r.service_name,
                r.subpackage_name,
                r.id,
                r.ip_or_domain,
                r.title,
                r.severity,
                r.plugin_output,
            ]
        )

    ws_system = wb.create_sheet("系统汇总")
    if system_summary:
        ws_system.append(list(system_summary[0].keys()))
        for row in system_summary:
            ws_system.append(list(row.values()))
    else:
        ws_system.append(["系统", "总漏洞", "IAP漏洞数", "非IAP漏洞数", "高危以上", "高危以上IAP", "漏洞种类数(标题去重)"])

    ws_system_service = wb.create_sheet("系统-服务汇总")
    if system_service_summary:
        ws_system_service.append(list(system_service_summary[0].keys()))
        for row in system_service_summary:
            ws_system_service.append(list(row.values()))
    else:
        ws_system_service.append(
            [
                "系统",
                "服务名",
                "总漏洞",
                "高危以上总漏洞",
                "漏洞种类数(标题去重)",
                "IAP总漏洞",
                "高危以上IAP总漏洞",
                "IAP漏洞种类数(标题去重)",
            ]
        )

    ws_system_iap_service_subpackage = wb.create_sheet("系统-IAP-服务-子包汇总")
    if system_iap_service_subpackage_summary:
        ws_system_iap_service_subpackage.append(list(system_iap_service_subpackage_summary[0].keys()))
        for row in system_iap_service_subpackage_summary:
            ws_system_iap_service_subpackage.append(list(row.values()))
    else:
        ws_system_iap_service_subpackage.append(
            ["系统", "服务名", "子包名称", "总漏洞", "高危以上", "漏洞种类数(标题去重)", "漏洞Id汇总", "漏洞标题汇总"]
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    """解析命令行参数。"""
    parser = argparse.ArgumentParser(description="漏洞整理分类：Excel 输入 -> 分类统计 Excel 输出")
    parser.add_argument("--input", required=True, help="输入 Excel 文件路径（.xlsx）")
    parser.add_argument("--output", required=True, help="输出 Excel 文件路径（.xlsx）")
    parser.add_argument("--config", required=True, help="配置文件路径（JSON）")
    parser.add_argument("--sheet", default="", help="可选：指定 Sheet 名称（默认使用活动页）")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    """程序入口：读取 Excel、分类统计并写出报表。"""
    args = parse_args(argv)
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    config_path = Path(args.config).expanduser().resolve()

    config = load_config(config_path)
    rows, _headers = read_vuln_records_from_excel(
        input_path=input_path,
        columns=config.columns,
        sheet_name=args.sheet.strip() or None,
    )

    records = build_records(rows, config)
    system_summary = summarize_by_system(records, config.high_levels)
    system_service_summary = summarize_by_system_service(records, config.high_levels)
    system_iap_service_subpackage_summary = summarize_by_system_iap_service_subpackage(records, config.high_levels)

    write_output_excel(
        output_path=output_path,
        records=records,
        system_summary=system_summary,
        system_service_summary=system_service_summary,
        system_iap_service_subpackage_summary=system_iap_service_subpackage_summary,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
